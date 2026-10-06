from __future__ import annotations

from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import (
    PageBreak,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)


def _text(value: Any) -> str:
    if value is None:
        return ""
    return str(value)


def _money(value: Any) -> str:
    try:
        return f"{float(value or 0):,.2f}"
    except (TypeError, ValueError):
        return _text(value)


def _statement_rows(statement: dict) -> list[list[str]]:
    rows: list[list[str]] = [["Code", "Account", "Amount"]]
    for section, items in (statement.get("sections") or {}).items():
        rows.append(["", str(section).replace("_", " ").title(), ""])
        for item in items or []:
            rows.append([
                _text(item.get("code")),
                _text(item.get("name")),
                _money(item.get("amount")),
            ])
        section_total = (statement.get("totals") or {}).get(section)
        if section_total is not None:
            rows.append(["", f"Total {str(section).replace('_', ' ').title()}", _money(section_total)])
    for key, value in (statement.get("totals") or {}).items():
        if key in (statement.get("sections") or {}):
            continue
        rows.append(["", str(key).replace("_", " ").title(), _money(value)])
    return rows


def _pdf_table(data: list[list[Any]], widths=None, *, header=True) -> Table:
    table = Table(data, colWidths=widths, repeatRows=1 if header else 0)
    style = [
        ("GRID", (0, 0), (-1, -1), 0.35, colors.grey),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#ECEFF1")),
        ("FONTSIZE", (0, 0), (-1, -1), 8),
        ("LEFTPADDING", (0, 0), (-1, -1), 4),
        ("RIGHTPADDING", (0, 0), (-1, -1), 4),
        ("TOPPADDING", (0, 0), (-1, -1), 3),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
    ]
    table.setStyle(TableStyle(style))
    return table


def build_financial_books_pdf(pack: dict, *, company_name: str) -> bytes:
    buffer = BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=landscape(A4),
        rightMargin=12 * mm,
        leftMargin=12 * mm,
        topMargin=14 * mm,
        bottomMargin=14 * mm,
        title=f"{company_name} Financial Books",
        author="LoanHub",
    )
    styles = getSampleStyleSheet()
    title = ParagraphStyle(
        "BookTitle",
        parent=styles["Title"],
        alignment=TA_CENTER,
        fontSize=20,
        leading=24,
        spaceAfter=8,
    )
    section = ParagraphStyle(
        "SectionTitle",
        parent=styles["Heading1"],
        fontSize=14,
        leading=17,
        spaceBefore=6,
        spaceAfter=8,
    )
    small = ParagraphStyle("Small", parent=styles["BodyText"], fontSize=8, leading=10)

    period = pack.get("period") or {}
    story = [
        Spacer(1, 18 * mm),
        Paragraph(company_name, title),
        Paragraph("LoanHub Financial Books", title),
        Paragraph(
            f"Reporting period: {_text(period.get('from_date'))} to {_text(period.get('to_date'))}",
            styles["Heading2"],
        ),
        Spacer(1, 8 * mm),
        Paragraph(_text(pack.get("accounting_basis")), styles["BodyText"]),
        Spacer(1, 4 * mm),
        Paragraph(_text(pack.get("preparation_note")), small),
        PageBreak(),
        Paragraph("Book Index", section),
    ]

    index_rows = [["#", "Book", "Purpose"]]
    for item in pack.get("book_index") or []:
        index_rows.append([
            _text(item.get("order")),
            _text(item.get("book")).replace("_", " ").title(),
            _text(item.get("purpose")),
        ])
    story += [_pdf_table(index_rows, [12 * mm, 55 * mm, 185 * mm]), PageBreak()]

    trial = pack.get("trial_balance") or {}
    story.append(Paragraph("Trial Balance", section))
    trial_rows = [["Code", "Account", "Type", "Debit", "Credit", "Balance"]]
    for line in trial.get("lines") or []:
        trial_rows.append([
            _text(line.get("code")),
            _text(line.get("name")),
            _text(line.get("account_type")),
            _money(line.get("debit")),
            _money(line.get("credit")),
            _money(line.get("balance")),
        ])
    trial_rows.append([
        "", "TOTAL", "",
        _money(trial.get("total_debit")),
        _money(trial.get("total_credit")),
        _money(trial.get("difference")),
    ])
    story += [_pdf_table(trial_rows, [22 * mm, 78 * mm, 32 * mm, 35 * mm, 35 * mm, 35 * mm]), PageBreak()]

    story.append(Paragraph("Income Statement", section))
    story += [_pdf_table(_statement_rows(pack.get("income_statement") or {}), [28 * mm, 170 * mm, 48 * mm]), PageBreak()]

    story.append(Paragraph("Statement of Financial Position", section))
    story += [_pdf_table(_statement_rows(pack.get("statement_of_financial_position") or {}), [28 * mm, 170 * mm, 48 * mm]), PageBreak()]

    story.append(Paragraph("Statement of Changes in Equity", section))
    equity_rows = [["Measure", "Amount"]]
    for key, value in (pack.get("statement_of_changes_in_equity") or {}).items():
        if isinstance(value, (dict, list)):
            continue
        equity_rows.append([str(key).replace("_", " ").title(), _money(value) if isinstance(value, (int, float)) else _text(value)])
    story += [_pdf_table(equity_rows, [150 * mm, 70 * mm]), PageBreak()]

    story.append(Paragraph("Statement of Cash Flows", section))
    cash = pack.get("statement_of_cash_flows") or {}
    cash_rows = [["Measure", "Value"]]
    for key, value in cash.items():
        if isinstance(value, (dict, list)):
            continue
        cash_rows.append([str(key).replace("_", " ").title(), _money(value) if isinstance(value, (int, float)) else _text(value)])
    story += [_pdf_table(cash_rows, [150 * mm, 70 * mm]), PageBreak()]

    story.append(Paragraph("General Ledger Summary", section))
    ledger_rows = [["Code", "Account", "Opening", "Debits", "Credits", "Closing"]]
    for row in pack.get("general_ledger") or []:
        ledger_rows.append([
            _text(row.get("account_code")),
            _text(row.get("account_name")),
            _money(row.get("opening_balance")),
            _money(row.get("period_debit")),
            _money(row.get("period_credit")),
            _money(row.get("closing_balance")),
        ])
    story += [_pdf_table(ledger_rows, [22 * mm, 92 * mm, 34 * mm, 34 * mm, 34 * mm, 34 * mm]), PageBreak()]

    story.append(Paragraph("Financial Ratios", section))
    ratio_rows = [["Group", "Ratio", "Value"]]
    for group, values in (pack.get("financial_ratios") or {}).items():
        if not isinstance(values, dict):
            continue
        for key, value in values.items():
            if isinstance(value, (dict, list)):
                continue
            ratio_rows.append([
                str(group).replace("_", " ").title(),
                str(key).replace("_", " ").title(),
                "" if value is None else _text(value),
            ])
    story += [_pdf_table(ratio_rows, [60 * mm, 130 * mm, 45 * mm]), PageBreak()]

    story.append(Paragraph("Accounting Controls", section))
    controls = ((pack.get("accounting_controls") or {}).get("error_diagnostics") or {}).get("controls") or {}
    control_rows = [["Control", "Status"]]
    for key, passed in controls.items():
        control_rows.append([str(key).replace("_", " ").title(), "PASS" if passed else "FAIL"])
    story += [_pdf_table(control_rows, [185 * mm, 40 * mm])]

    doc.build(story)
    return buffer.getvalue()


def _setup_sheet(ws, title: str) -> None:
    ws.title = title[:31]
    ws.freeze_panes = "A2"
    ws.sheet_view.showGridLines = False


def _append_header(ws, values: list[Any]) -> None:
    ws.append(values)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center")


def _auto_width(ws) -> None:
    for column in ws.columns:
        letter = column[0].column_letter
        max_length = 0
        for cell in column:
            value = "" if cell.value is None else str(cell.value)
            max_length = max(max_length, min(len(value), 50))
        ws.column_dimensions[letter].width = max(12, max_length + 2)


def build_financial_books_xlsx(pack: dict, *, company_name: str) -> bytes:
    wb = Workbook()
    index = wb.active
    _setup_sheet(index, "Index")
    index.append([company_name])
    index["A1"].font = Font(bold=True, size=16)
    period = pack.get("period") or {}
    index.append(["Reporting period", period.get("from_date"), period.get("to_date")])
    index.append(["Accounting basis", pack.get("accounting_basis")])
    index.append([])
    index.append(["#", "Book", "Purpose"])
    for item in pack.get("book_index") or []:
        index.append([item.get("order"), item.get("book"), item.get("purpose")])

    tb = wb.create_sheet()
    _setup_sheet(tb, "Trial Balance")
    _append_header(tb, ["Code", "Account", "Type", "Debit", "Credit", "Balance"])
    for line in (pack.get("trial_balance") or {}).get("lines") or []:
        tb.append([line.get("code"), line.get("name"), line.get("account_type"), line.get("debit"), line.get("credit"), line.get("balance")])

    for sheet_name, key in [
        ("Income Statement", "income_statement"),
        ("Financial Position", "statement_of_financial_position"),
    ]:
        ws = wb.create_sheet()
        _setup_sheet(ws, sheet_name)
        _append_header(ws, ["Section", "Code", "Account", "Amount"])
        statement = pack.get(key) or {}
        for section_name, rows in (statement.get("sections") or {}).items():
            for row in rows or []:
                ws.append([section_name, row.get("code"), row.get("name"), row.get("amount")])
        for total_name, total_value in (statement.get("totals") or {}).items():
            ws.append(["TOTAL", total_name, "", total_value])

    equity = wb.create_sheet()
    _setup_sheet(equity, "Changes in Equity")
    _append_header(equity, ["Measure", "Value"])
    for key, value in (pack.get("statement_of_changes_in_equity") or {}).items():
        if not isinstance(value, (dict, list)):
            equity.append([key, value])

    cash = wb.create_sheet()
    _setup_sheet(cash, "Cash Flow")
    _append_header(cash, ["Measure", "Value"])
    for key, value in (pack.get("statement_of_cash_flows") or {}).items():
        if not isinstance(value, (dict, list)):
            cash.append([key, value])

    ledger = wb.create_sheet()
    _setup_sheet(ledger, "General Ledger")
    _append_header(ledger, [
        "Account Code", "Account", "Opening", "Entry Date", "Entry Number",
        "Description", "Reference Type", "Reference ID", "Debit", "Credit",
        "Running Balance", "Closing",
    ])
    for account in pack.get("general_ledger") or []:
        lines = account.get("lines") or []
        if not lines:
            ledger.append([
                account.get("account_code"), account.get("account_name"),
                account.get("opening_balance"), "", "", "", "", "", "",
                "", "", account.get("closing_balance"),
            ])
            continue
        for line in lines:
            ledger.append([
                account.get("account_code"), account.get("account_name"),
                account.get("opening_balance"), line.get("entry_date"),
                line.get("entry_number"), line.get("description"),
                line.get("reference_type"), line.get("reference_id"),
                line.get("debit"), line.get("credit"),
                line.get("running_balance"), account.get("closing_balance"),
            ])

    original = wb.create_sheet()
    _setup_sheet(original, "Original Entry Books")
    _append_header(original, [
        "Date", "Book", "Entry", "Narrative", "Reference Type",
        "Reference ID", "Debit", "Credit",
    ])
    for row in pack.get("books_of_original_entry") or []:
        original.append([
            row.get("entry_date"), row.get("book"), row.get("entry_number"),
            row.get("narrative"), row.get("reference_type"), row.get("reference_id"),
            row.get("total_debit"), row.get("total_credit"),
        ])

    ratios = wb.create_sheet()
    _setup_sheet(ratios, "Ratios")
    _append_header(ratios, ["Group", "Ratio", "Value"])
    for group, values in (pack.get("financial_ratios") or {}).items():
        if not isinstance(values, dict):
            continue
        for key, value in values.items():
            if not isinstance(value, (dict, list)):
                ratios.append([group, key, value])

    controls = wb.create_sheet()
    _setup_sheet(controls, "Controls")
    _append_header(controls, ["Control", "Status"])
    error_controls = ((pack.get("accounting_controls") or {}).get("error_diagnostics") or {}).get("controls") or {}
    for key, passed in error_controls.items():
        controls.append([key, "PASS" if passed else "FAIL"])

    for ws in wb.worksheets:
        _auto_width(ws)
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, float):
                    cell.number_format = '#,##0.00'

    output = BytesIO()
    wb.save(output)
    return output.getvalue()
