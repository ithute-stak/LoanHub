from __future__ import annotations

import base64
import csv
import hashlib
import io
import re
from datetime import date, datetime, time
from decimal import Decimal, InvalidOperation
from typing import Any
from zoneinfo import ZoneInfo

from openpyxl import load_workbook
from sqlalchemy.orm import Session

from database.models.origination import OriginationIntegrationConfiguration
from database.models.platform_cdas import PlatformCdasTransaction
from integrations.cdas import CdasDocumentType, CdasError
from services.cdas_config_service import (
    CDAS_PROVIDER,
    get_company_cdas_client,
    selected_environment,
)
from services.platform_cdas_service import (
    assert_live_credit_available,
    record_successful_operation,
    require_approved_subscription,
)


CDAS_TIMEZONE = ZoneInfo("Africa/Maseru")
DAILY_SYNC_TIME = time(hour=3, minute=45)
MAX_SNAPSHOT_EMPLOYEES = 5000

_EMPLOYEE_HEADERS = {
    "employeeno", "employeenumber", "employeeid", "staffno", "staffnumber",
    "payrollno", "payrollnumber", "personnelno", "personnelnumber",
}
_NAME_HEADERS = {"clientname", "employeename", "fullname", "name"}
_FIRST_HEADERS = {"firstname", "firstnames", "givenname", "givennames"}
_SURNAME_HEADERS = {"surname", "lastname", "familyname"}
_NID_HEADERS = {"nationalid", "nationalidnumber", "idnumber", "idno", "nid"}
_EMPLOYER_HEADERS = {"employer", "employername", "department", "ministry"}
_AMOUNT_HEADERS = {
    "amount", "deductionamount", "monthlydeduction", "monthlydeductionamount",
    "collectionamount", "installmentamount", "instalmentamount",
}
_STATUS_HEADERS = {"status", "deductionstatus"}
_REFERENCE_HEADERS = {
    "reference", "referenceno", "referencenumber", "deductionreference", "deductionid",
}


class CdasRosterDocumentError(ValueError):
    pass


def _normalize_header(value: Any) -> str:
    return re.sub(r"[^a-z0-9]", "", str(value or "").strip().lower())


def _clean(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _first(row: dict[str, Any], aliases: set[str]) -> str:
    for key, value in row.items():
        if _normalize_header(key) in aliases:
            result = _clean(value)
            if result:
                return result
    return ""


def _decimal(value: Any) -> Decimal | None:
    text = _clean(value).replace(",", "")
    if not text:
        return None
    match = re.search(r"-?\d+(?:\.\d+)?", text)
    if not match:
        return None
    try:
        return Decimal(match.group(0))
    except InvalidOperation:
        return None


def _decode_content(document: dict[str, Any]) -> bytes:
    value = document.get("Content", document.get("content"))
    if value is None:
        raise CdasRosterDocumentError("CDAS output file did not contain Content")
    if isinstance(value, bytes):
        return value
    if isinstance(value, bytearray):
        return bytes(value)
    if isinstance(value, list) and all(isinstance(item, int) and 0 <= item <= 255 for item in value):
        return bytes(value)
    if isinstance(value, dict):
        raw = value.get("data", value.get("Data"))
        if isinstance(raw, list) and all(isinstance(item, int) and 0 <= item <= 255 for item in raw):
            return bytes(raw)
        raise CdasRosterDocumentError("CDAS output file used an unsupported Content object")
    if not isinstance(value, str):
        raise CdasRosterDocumentError("CDAS output file used an unsupported Content format")
    text = value.strip()
    if text.startswith("data:") and "," in text:
        header, text = text.split(",", 1)
        if ";base64" not in header.lower():
            return text.encode("utf-8")
    compact = re.sub(r"\s+", "", text)
    if len(compact) >= 16:
        try:
            decoded = base64.b64decode(compact + ("=" * (-len(compact) % 4)), validate=True)
            if decoded.startswith(b"PK\x03\x04") or b"\n" in decoded:
                return decoded
        except Exception:
            pass
    return text.encode("utf-8")


def _header_row(matrix: list[list[Any]]) -> tuple[int, list[str]]:
    for index, values in enumerate(matrix[:25]):
        headers = [_normalize_header(value) for value in values]
        if any(value in _EMPLOYEE_HEADERS for value in headers):
            return index, [_clean(value) for value in values]
    raise CdasRosterDocumentError(
        "CDAS output file has no recognised employee-number header; LoanHub refused to guess columns"
    )


def _xlsx_rows(data: bytes) -> list[dict[str, Any]]:
    try:
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:
        raise CdasRosterDocumentError("CDAS output file could not be opened as XLSX") from exc
    try:
        for sheet in workbook.worksheets:
            matrix = [list(row) for row in sheet.iter_rows(values_only=True)]
            if not matrix:
                continue
            try:
                header_index, headers = _header_row(matrix)
            except CdasRosterDocumentError:
                continue
            return [
                {headers[i]: value for i, value in enumerate(row) if i < len(headers) and headers[i]}
                for row in matrix[header_index + 1:]
                if any(_clean(value) for value in row)
            ]
    finally:
        workbook.close()
    raise CdasRosterDocumentError("No worksheet contained a recognised employee-number header")


def _text_rows(data: bytes) -> list[dict[str, Any]]:
    try:
        text = data.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = data.decode("cp1252")
    try:
        dialect = csv.Sniffer().sniff(text[:8192], delimiters=",;\t|")
    except csv.Error as exc:
        raise CdasRosterDocumentError("CDAS output delimiter could not be identified safely") from exc
    matrix = [list(row) for row in csv.reader(io.StringIO(text), dialect)]
    header_index, headers = _header_row(matrix)
    return [
        {headers[i]: value for i, value in enumerate(row) if i < len(headers) and headers[i]}
        for row in matrix[header_index + 1:]
        if any(_clean(value) for value in row)
    ]


def parse_cdas_roster_document(document: dict[str, Any]) -> tuple[list[dict[str, Any]], str]:
    data = _decode_content(document)
    filename = _clean(document.get("FileName") or document.get("fileName") or document.get("filename"))
    lower = filename.lower()
    if data.startswith(b"PK\x03\x04") or lower.endswith((".xlsx", ".xlsm")):
        rows = _xlsx_rows(data)
    elif lower.endswith(".xls"):
        raise CdasRosterDocumentError("Legacy .xls is not parsed because column safety cannot be guaranteed")
    else:
        rows = _text_rows(data)
    return rows, hashlib.sha256(data).hexdigest()


def build_roster_snapshot(rows: list[dict[str, Any]]) -> dict[str, Any]:
    grouped: dict[str, dict[str, Any]] = {}
    for row in rows:
        employee_no = _first(row, _EMPLOYEE_HEADERS)
        if not employee_no:
            continue
        item = grouped.setdefault(
            employee_no,
            {
                "employee_number": employee_no,
                "name": "",
                "national_id": "",
                "employer": "",
                "row_count": 0,
                "monthly_deductions": Decimal("0.00"),
                "amount_values": 0,
                "references": set(),
                "statuses": set(),
            },
        )
        full_name = _first(row, _NAME_HEADERS)
        if not full_name:
            full_name = " ".join(
                value for value in (_first(row, _FIRST_HEADERS), _first(row, _SURNAME_HEADERS)) if value
            ).strip()
        if full_name:
            item["name"] = full_name
        national_id = _first(row, _NID_HEADERS)
        if national_id:
            item["national_id"] = national_id
        employer = _first(row, _EMPLOYER_HEADERS)
        if employer:
            item["employer"] = employer
        item["row_count"] += 1
        amount = _decimal(_first(row, _AMOUNT_HEADERS))
        if amount is not None:
            item["monthly_deductions"] += amount
            item["amount_values"] += 1
        reference = _first(row, _REFERENCE_HEADERS)
        if reference:
            item["references"].add(reference)
        status = _first(row, _STATUS_HEADERS)
        if status:
            item["statuses"].add(status)

    normalized = []
    total_deductions = Decimal("0.00")
    for employee_no in sorted(grouped):
        item = grouped[employee_no]
        total_deductions += item["monthly_deductions"]
        normalized.append(
            {
                "employee_number": item["employee_number"],
                "name": item["name"] or None,
                "national_id": item["national_id"] or None,
                "employer": item["employer"] or None,
                "row_count": item["row_count"],
                "monthly_deductions": str(item["monthly_deductions"].quantize(Decimal("0.01"))),
                "references": sorted(item["references"]),
                "statuses": sorted(item["statuses"]),
                "data_quality": "complete" if item["amount_values"] else "amount_missing",
            }
        )

    return {
        "employee_count": len(normalized),
        "total_monthly_deductions": str(total_deductions.quantize(Decimal("0.01"))),
        "employees": normalized[:MAX_SNAPSHOT_EMPLOYEES],
        "truncated": len(normalized) > MAX_SNAPSHOT_EMPLOYEES,
    }


def _month_target(now: datetime) -> tuple[int, int]:
    return now.year, now.month


def _billing_key(company_id, environment: str, day: date) -> str:
    return f"cdas-roster:{company_id}:{environment}:{day.isoformat()}"


def latest_roster_snapshot(db: Session, *, company_id) -> dict[str, Any] | None:
    rows = (
        db.query(PlatformCdasTransaction)
        .filter(
            PlatformCdasTransaction.company_id == company_id,
            PlatformCdasTransaction.operation_type == "document",
        )
        .order_by(PlatformCdasTransaction.accrued_at.desc())
        .limit(100)
        .all()
    )
    for row in rows:
        metadata = dict(row.metadata_json or {})
        if metadata.get("snapshot_kind") == "cdas_roster":
            return {
                "transaction_id": str(row.id),
                "environment": row.environment,
                "captured_at": row.accrued_at,
                "source_reference": row.source_reference,
                **metadata,
            }
    return None


async def sync_company_roster(
    db: Session,
    *,
    company_id,
    actor_user_id=None,
    now: datetime | None = None,
) -> dict[str, Any]:
    current = now.astimezone(CDAS_TIMEZONE) if now else datetime.now(CDAS_TIMEZONE)
    environment_row = (
        db.query(OriginationIntegrationConfiguration)
        .filter(
            OriginationIntegrationConfiguration.company_id == company_id,
            OriginationIntegrationConfiguration.provider == CDAS_PROVIDER,
        )
        .one_or_none()
    )
    environment = selected_environment(environment_row)
    key = _billing_key(company_id, environment, current.date())
    existing = (
        db.query(PlatformCdasTransaction)
        .filter(PlatformCdasTransaction.billing_key == key)
        .first()
    )
    if existing:
        metadata = dict(existing.metadata_json or {})
        return {"ok": True, "reused": True, "transaction_id": str(existing.id), **metadata}

    subscription = require_approved_subscription(db, company_id=company_id)
    assert_live_credit_available(
        db,
        subscription=subscription,
        environment=environment,
        operation_type="document",
    )
    client = get_company_cdas_client(db, company_id)
    year, month = _month_target(current)
    document = await client.get_document(
        year=year,
        month=month,
        document_type=int(CdasDocumentType.OUTPUT_FILE),
    )
    rows, content_sha256 = parse_cdas_roster_document(document)
    snapshot = build_roster_snapshot(rows)
    filename = _clean(document.get("FileName") or document.get("fileName") or document.get("filename"))
    transaction = record_successful_operation(
        db,
        company_id=company_id,
        environment=environment,
        operation_type="document",
        actor_user_id=actor_user_id,
        billing_key=key,
        source_reference=f"roster:{year:04d}-{month:02d}",
        metadata={
            "request_origin": "cdas_roster_intelligence",
            "snapshot_kind": "cdas_roster",
            "snapshot_date": current.date().isoformat(),
            "provider_year": year,
            "provider_month": month,
            "file_name": filename or None,
            "content_sha256": content_sha256,
            **snapshot,
        },
    )
    return {
        "ok": True,
        "reused": False,
        "transaction_id": str(transaction.id),
        **dict(transaction.metadata_json or {}),
    }


def roster_sync_due(db: Session, *, company_id, now: datetime) -> bool:
    local = now.astimezone(CDAS_TIMEZONE)
    if local.timetz().replace(tzinfo=None) < DAILY_SYNC_TIME:
        return False
    environment_row = (
        db.query(OriginationIntegrationConfiguration)
        .filter(
            OriginationIntegrationConfiguration.company_id == company_id,
            OriginationIntegrationConfiguration.provider == CDAS_PROVIDER,
        )
        .one_or_none()
    )
    environment = selected_environment(environment_row)
    return (
        db.query(PlatformCdasTransaction)
        .filter(PlatformCdasTransaction.billing_key == _billing_key(company_id, environment, local.date()))
        .first()
        is None
    )


async def run_roster_sync_cycle(db: Session, *, now: datetime | None = None) -> list[dict[str, Any]]:
    current = now or datetime.now(CDAS_TIMEZONE)
    configurations = (
        db.query(OriginationIntegrationConfiguration)
        .filter(
            OriginationIntegrationConfiguration.provider == CDAS_PROVIDER,
            OriginationIntegrationConfiguration.is_enabled.is_(True),
        )
        .all()
    )
    results = []
    for row in configurations:
        if not roster_sync_due(db, company_id=row.company_id, now=current):
            continue
        try:
            results.append(await sync_company_roster(db, company_id=row.company_id, now=current))
        except (CdasError, CdasRosterDocumentError, ValueError) as exc:
            db.rollback()
            results.append({"ok": False, "company_id": str(row.company_id), "error": str(exc)})
        except Exception as exc:
            db.rollback()
            results.append({"ok": False, "company_id": str(row.company_id), "error": str(exc)})
    return results
