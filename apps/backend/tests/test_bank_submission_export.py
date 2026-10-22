import pytest
"""Regression tests for per-bank CSV account widths and text-safe Excel exports."""
import csv
from io import BytesIO, StringIO
from zipfile import ZipFile

from openpyxl import load_workbook
from services.bank_submission_export import build_bank_submission_zip


def _record(**overrides):
    row = dict(
        name="Thato M", account="60123456789", branch="280061",
        bank="FNB", group="FORCE", installment="1000", balance="2400",
    )
    row.update(overrides)
    return row


def _read(archive, name):
    return list(csv.reader(StringIO(archive.read(name).decode("utf-8"))))


def test_bank_account_lengths_branch_and_name_format():
    records = [
        _record(name="Ntai Lephe", account="601234567", branch="280061"),
        _record(name="Pekechu Lephe", bank="STD", account="9080000368791", branch="60667", group="CIVIL"),
        _record(name="Mathabo Se", bank="PB", account="10123456789", branch="500100"),
        _record(name="Lefusi M", bank="NB", account="123456789", branch="390161"),
    ]
    payload, stats = build_bank_submission_zip(records, lender_name="Batlokoa")
    with ZipFile(BytesIO(payload)) as archive:
        rows = [row for name in archive.namelist() for row in _read(archive, name)]
        assert ["Lephe N", "00601234567", "280061", "1000.00", "Batlokoa"] in rows
        assert ["Lephe P", "9080000368791", "060667", "1000.00", "Batlokoa"] in rows
        assert ["Se M", "0010123456789", "500100", "1000.00", "Batlokoa"] in rows
        assert ["Lefusi M", "00123456789", "390161", "1000.00", "Batlokoa"] in rows
    assert stats["rows"] == 4
    assert stats["excluded"] == {}


def test_existing_surname_initial_not_swapped():
    payload, _ = build_bank_submission_zip([_record(name="Mosipatsi P")], lender_name="Batlokoa")
    with ZipFile(BytesIO(payload)) as archive:
        assert _read(archive, "FNB-all-groups.csv")[0][0] == "Mosipats P"


def test_invalid_or_overlength_accounts_are_rejected():
    invalid = [
        _record(bank="STD", account="908000036879100"),
        _record(bank="FNB", account="601234567890"),
        _record(bank="NB", account="123456789012"),
        _record(bank="PB", account="10123456789123"),
        _record(account="****4567"),
        _record(branch="1234567"),
        _record(balance="0"),
    ]
    payload, stats = build_bank_submission_zip(invalid, lender_name="Batlokoa")
    with ZipFile(BytesIO(payload)) as archive:
        assert not archive.namelist()
    assert stats["rows"] == 0
    assert sum(stats["excluded"].values()) == len(invalid)


def test_formatted_spreadsheet_companion_preserves_identifiers():
    payload, stats = build_bank_submission_zip(
        [_record(account="601234567", branch="61")],
        lender_name="Batlokoa",
        include_spreadsheets=True,
    )
    with ZipFile(BytesIO(payload)) as archive:
        row = _read(archive, "FNB-all-groups.csv")[0]
        assert row[1:3] == ["00601234567", "000061"]
        workbook = load_workbook(BytesIO(archive.read("FNB-all-groups.xlsx")))
        sheet = workbook.active
        assert sheet["B1"].value == "00601234567"
        assert sheet["C1"].value == "000061"
        assert sheet["B1"].data_type == "s"
        assert sheet["C1"].data_type == "s"
        assert sheet["B1"].number_format == "@"
        assert sheet["C1"].number_format == "@"
        assert sheet.column_dimensions["B"].width >= 20
        assert sheet.column_dimensions["C"].width >= 12
    assert stats["rows"] == 1
    assert stats["files"] == 1

def test_rejected_indices_track_original_record_order():
    payload, stats = build_bank_submission_zip(
        [
            _record(name="Valid Client", account="60123456789"),
            _record(name="Invalid Client", account="*****"),
            _record(name="Valid Other", account="60123456788"),
        ],
        lender_name="Batlokoa",
    )
    assert stats["rows"] == 2
    assert stats["rejected_indices"] == [
        {"index": 1, "reason": "missing_or_invalid_bank_details"},
    ]
    with ZipFile(BytesIO(payload)) as archive:
        assert len(_read(archive, "FNB-all-groups.csv")) == 2

def test_zip_includes_csv_import_guidance_without_mutating_bank_data():
    payload, _ = build_bank_submission_zip(
        [_record(account="60123456789", branch="280061")],
        lender_name="Batlokoa",
        include_spreadsheets=True,
    )
    with ZipFile(BytesIO(payload)) as archive:
        assert "READ-ME-FIRST-CSV-AND-EXCEL.txt" in archive.namelist()
        guide = archive.read("READ-ME-FIRST-CSV-AND-EXCEL.txt").decode("utf-8")
        assert "columns B and C as TEXT" in guide
        assert _read(archive, "FNB-all-groups.csv")[0][1:3] == [
            "60123456789", "280061",
        ]


def test_csv_amount_is_unquoted_machine_readable_decimal():
    payload, _ = build_bank_submission_zip(
        [_record(name="Lephe N", account="60123456789", installment="125.5",
                 balance="500")], lender_name="Batlokoa",
    )
    with ZipFile(BytesIO(payload)) as archive:
        raw = archive.read("FNB-all-groups.csv").decode("utf-8").strip()
        assert raw == "Lephe N,60123456789,280061,125.50,Batlokoa"
        assert _read(archive, "FNB-all-groups.csv")[0][3] == "125.50"


@pytest.mark.parametrize("bank,account,branch", [
    ("FNB", "60123456789", "280061"),
    ("STD", "9012345678901", "060667"),
    ("PB", "1012345678901", "500100"),
    ("NB", "12345678901", "390161"),
])
def test_bank_csv_matches_accepted_windows_samples_without_customer_data(bank, account, branch):
    # Structural properties measured from seven recipient-accepted CSV files:
    # five fields, no header/BOM/quotes, CRLF, 6-digit branch, two decimal places.
    payload, result = build_bank_submission_zip(
        [_record(name="Lephe N", bank=bank, account=account,
                 branch=branch, installment="125.5", balance="500")],
        lender_name="Batlokoa", include_spreadsheets=True,
    )
    assert result["rows"] == 1
    with ZipFile(BytesIO(payload)) as archive:
        csv_name = next(name for name in archive.namelist() if name.endswith(".csv"))
        raw = archive.read(csv_name)
        assert not raw.startswith(b"\xef\xbb\xbf")
        assert raw.endswith(b"\r\n") and raw.count(b"\n") == 1
        assert b'"' not in raw
        assert raw == f"Lephe N,{account},{branch},125.50,Batlokoa\r\n".encode("utf-8")
        row = next(csv.reader(StringIO(raw.decode("utf-8"))))
        assert len(row) == 5
        assert len(row[1]) == (11 if bank in {"FNB", "NB"} else 13)
        assert len(row[2]) == 6
        workbook = load_workbook(BytesIO(archive.read(csv_name[:-4] + ".xlsx")))
        sheet = workbook.active
        assert sheet["B1"].value == account and sheet["B1"].data_type == "s"
        assert sheet["C1"].value == branch and sheet["C1"].data_type == "s"
        assert sheet["D1"].value == 125.5 and sheet["D1"].data_type == "n"
        assert sheet["D1"].number_format == "0.00"
